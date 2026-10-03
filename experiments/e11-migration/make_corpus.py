#!/usr/bin/env python3
"""E11 corpus: a synthetic SharePoint library of Word documents and a synthetic
Confluence space, both about a fictional company's payments platform
("Northwind Pay"). Every element that a migration could keep or lose is
recorded in corpus/ground-truth.json with a probe string, so the measurement
step can ask of each output: did this survive, and in what form?

Usage: python3 make_corpus.py <out-dir>      (needs python-docx >= 1.2)
"""
from __future__ import annotations

import json
import struct
import sys
import zlib
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "corpus")
SITE = "https://example.sharepoint.com/sites/PaymentsPlatform"
LIB = "Shared Documents"
WIKI = "https://example.atlassian.net/wiki"
truth: list[dict] = []


def probe(source: str, doc: str, kind: str, text: str, **extra) -> None:
    truth.append({"source": source, "doc": doc, "kind": kind, "probe": text, **extra})


def png(colour: tuple[int, int, int]) -> bytes:
    """A 16x16 single-colour PNG, built by hand so no image library is needed."""
    w = h = 16
    raw = b"".join(b"\x00" + bytes(colour) * w for _ in range(h))
    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


# ---------------------------------------------------------------- Word / SharePoint

def doc_url(folder: str, name: str) -> str:
    return f"{SITE}/{LIB.replace(' ', '%20')}/{folder}/{name.replace(' ', '%20')}.docx"


WORD = [
    {"file": "Payment Service Overview", "folder": "Designs", "docid": "NWP-1001", "ctype": "Design",
     "author": "Jane Doe", "modby": "Sam Patel", "created": "2025-11-04T09:00:00Z", "modified": "2026-07-18T09:12:00Z",
     "approval": {"status": "Approved", "by": "Alex Kim", "at": "2026-07-18T11:00:00Z"},
     "terms": ["Payments", "Architecture"], "subject": "Payment service design", "keywords": "payments; capture; design"},
    {"file": "Capture Runbook", "folder": "Runbooks", "docid": "NWP-1002", "ctype": "Runbook",
     "author": "Sam Patel", "modby": "Sam Patel", "created": "2026-01-12T10:00:00Z", "modified": "2026-08-30T14:20:00Z",
     "approval": {"status": "Approved", "by": "Jane Doe", "at": "2026-09-01T08:30:00Z"},
     "terms": ["Payments", "Operations"], "subject": "Capture runbook", "keywords": "capture; runbook; on-call"},
    {"file": "Refund Policy", "folder": "Policies", "docid": "NWP-1003", "ctype": "Policy",
     "author": "Alex Kim", "modby": "Alex Kim", "created": "2026-05-02T13:00:00Z", "modified": "2026-09-20T16:45:00Z",
     "approval": {"status": "Pending"}, "terms": ["Payments", "Finance"], "subject": "Refund policy",
     "keywords": "refunds; policy"},
    {"file": "Orders Table Data Dictionary", "folder": "Data", "docid": "NWP-1004", "ctype": "DataAsset",
     "author": "Jane Doe", "modby": "Jane Doe", "created": "2025-03-10T09:30:00Z", "modified": "2025-08-14T10:00:00Z",
     "approval": {"status": "Approved", "by": "Sam Patel", "at": "2025-06-02T09:00:00Z"},
     "terms": ["Data", "Orders"], "subject": "Orders table", "keywords": "orders; schema; data dictionary"},
    {"file": "Incident Review Capture Latency August", "folder": "Reviews", "docid": "NWP-1005", "ctype": "Document",
     "author": "Alex Kim", "modby": "Jane Doe", "created": "2026-08-22T08:00:00Z", "modified": "2026-08-25T17:30:00Z",
     "approval": {"status": "Draft"}, "terms": ["Operations", "Incidents"], "subject": "Incident review",
     "keywords": "incident; latency; capture"},
]


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def tracked(paragraph, kind: str, text: str, author: str, when: str) -> None:
    """Append a tracked insertion (w:ins) or deletion (w:del) to a paragraph."""
    el = OxmlElement(f"w:{kind}")
    el.set(qn("w:id"), str(len(truth) + 100))
    el.set(qn("w:author"), author)
    el.set(qn("w:date"), when)
    run = OxmlElement("w:r")
    t = OxmlElement("w:delText" if kind == "del" else "w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    el.append(run)
    paragraph._p.append(el)


def word_doc(meta: dict, build) -> None:
    d = Document()
    cp = d.core_properties
    cp.title, cp.author, cp.last_modified_by = meta["file"], meta["author"], meta["modby"]
    cp.subject, cp.keywords, cp.revision = meta["subject"], meta["keywords"], 7
    cp.created = datetime.fromisoformat(meta["created"].replace("Z", "+00:00"))
    cp.modified = datetime.fromisoformat(meta["modified"].replace("Z", "+00:00"))
    name = meta["file"]
    for field, key in (("title", "file"), ("author", "author"), ("last-modified-by", "modby"), ("created", "created"),
                       ("modified", "modified"), ("subject", "subject"), ("keywords", "keywords"),
                       ("document-id", "docid"), ("content-type", "ctype")):
        probe("word", name, "metadata", str(meta[key]), field=field)
    probe("word", name, "metadata", meta["approval"]["status"], field="approval")
    for t in meta["terms"]:
        probe("word", name, "metadata", t, field="managed-metadata")
    build(d, name)
    path = OUT / "word" / meta["folder"]
    path.mkdir(parents=True, exist_ok=True)
    d.save(path / f"{name}.docx")
    history = [{"version": f"{v}.0", "modified": meta["created"] if v == 1 else meta["modified"],
                "modifiedBy": meta["author"] if v == 1 else meta["modby"],
                "comment": "Initial upload" if v == 1 else f"Revision {v}"} for v in (1, 2, 3)]
    history[1]["modified"] = meta["created"][:10] + "T15:00:00Z"
    for h in history:
        probe("word", name, "version", h["comment"], version=h["version"])
    sidecar = {"site": SITE, "library": LIB, "folder": meta["folder"], "url": doc_url(meta["folder"], name),
               "documentId": meta["docid"], "contentType": meta["ctype"], "managedMetadata": meta["terms"],
               "approval": meta["approval"], "versions": history,
               "permissions": "inherited from library"}
    (path / f"{name}.sharepoint.json").write_text(json.dumps(sidecar, indent=2), encoding="utf-8")


def heading(d, name, text, level=1):
    d.add_heading(text, level=level)
    probe("word", name, "heading", text)


def para(d, text):
    return d.add_paragraph(text)


def numbered(d, name, items):
    for it in items:
        d.add_paragraph(it, style="List Number")
        probe("word", name, "list-item", it)


def table(d, name, rows):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    probe("word", name, "table", rows[-1][0])


def link(d, name, url, text, internal_to=None):
    p = d.add_paragraph("See ")
    add_hyperlink(p, url, text)
    probe("word", name, "link-internal" if internal_to else "link-external", url,
          anchor=text, target=internal_to)


def comment(d, name, paragraph, text, author):
    d.add_comment(paragraph.runs[:1] or [paragraph.add_run(" ")], text=text, author=author,
                  initials="".join(w[0] for w in author.split()))
    probe("word", name, "comment", text, author=author)


def build_overview(d, n):
    heading(d, n, "Payment Service Overview", 0)
    p = para(d, "The payment service authorises and captures customer payments for Northwind Pay.")
    comment(d, n, p, "Should we say which card schemes are in scope?", "Alex Kim")
    heading(d, n, "Responsibilities")
    numbered(d, n, ["Authorise the payment with the acquirer", "Capture the funds after dispatch",
                    "Publish a settlement event"])
    heading(d, n, "Capacity")
    table(d, n, [["Region", "Peak TPS"], ["Europe", "1,200"], ["North America", "1,850"]])
    p = para(d, "Capture currently runs synchronously.")
    tracked(p, "ins", " It moves to an event-driven design in v2.", "Sam Patel", "2026-07-18T08:00:00Z")
    probe("word", n, "tracked-insert", "It moves to an event-driven design in v2.")
    tracked(p, "del", " Batch capture is planned for 2027.", "Sam Patel", "2026-07-18T08:00:00Z")
    probe("word", n, "tracked-delete", "Batch capture is planned for 2027.")
    heading(d, n, "Architecture")
    img = OUT / "_tmp-flow.png"
    img.parent.mkdir(parents=True, exist_ok=True)
    img.write_bytes(png((40, 90, 160)))
    d.add_picture(str(img), width=Inches(1.0))
    img.unlink()
    probe("word", n, "image", "payment-flow")
    link(d, n, doc_url("Runbooks", "Capture Runbook"), "the capture runbook", internal_to="Capture Runbook")
    link(d, n, "https://www.pcisecuritystandards.org/", "PCI DSS")


def build_runbook(d, n):
    heading(d, n, "Capture Runbook", 0)
    para(d, "Use this runbook when capture latency rises above the alert threshold.")
    heading(d, n, "Steps")
    numbered(d, n, ["Check the acquirer status page", "Drain the capture queue to the standby worker",
                    "Page the payments on-call engineer"])
    p = para(d, "Escalate after fifteen minutes.")
    comment(d, n, p, "Fifteen minutes is too long for peak trading.", "Jane Doe")
    link(d, n, doc_url("Designs", "Payment Service Overview"), "the payment service overview",
         internal_to="Payment Service Overview")


def build_refund(d, n):
    heading(d, n, "Refund Policy", 0)
    para(d, "Refunds return funds to the original payment method.")
    heading(d, n, "Limits")
    table(d, n, [["Refund type", "Approval needed"], ["Under 500 GBP", "None"], ["Over 500 GBP", "Finance lead"]])
    p = para(d, "Refunds are processed within five working days.")
    tracked(p, "del", " Refunds older than one year are refused.", "Alex Kim", "2026-09-20T16:00:00Z")
    probe("word", n, "tracked-delete", "Refunds older than one year are refused.")
    tracked(p, "ins", " Refunds older than two years need a finance lead.", "Alex Kim", "2026-09-20T16:00:00Z")
    probe("word", n, "tracked-insert", "Refunds older than two years need a finance lead.")
    link(d, n, doc_url("Designs", "Payment Service Overview"), "payment service", internal_to="Payment Service Overview")


def build_orders(d, n):
    heading(d, n, "Orders Table Data Dictionary", 0)
    para(d, "One row per completed customer order.")
    heading(d, n, "Columns")
    table(d, n, [["Column", "Type", "Meaning"], ["order_id", "uuid", "Primary key"],
                 ["captured_amount", "numeric", "Amount captured in minor units"]])
    link(d, n, doc_url("Designs", "Payment Service Overview"), "the payment service", internal_to="Payment Service Overview")
    link(d, n, "https://www.postgresql.org/docs/current/datatype-numeric.html", "PostgreSQL numeric types")


def build_incident(d, n):
    heading(d, n, "Incident Review: Capture Latency, August", 0)
    p = para(d, "On 21 August capture latency rose to nine seconds for forty minutes.")
    comment(d, n, p, "Was the acquirer at fault or our queue?", "Sam Patel")
    heading(d, n, "Timeline")
    numbered(d, n, ["14:05 alert fired", "14:20 queue drained to standby", "14:45 latency back to normal"])
    heading(d, n, "Actions")
    p = para(d, "Lower the escalation time in the runbook.")
    comment(d, n, p, "Agreed, ten minutes.", "Jane Doe")
    link(d, n, doc_url("Runbooks", "Capture Runbook"), "capture runbook", internal_to="Capture Runbook")
    link(d, n, doc_url("Designs", "Payment Service Overview"), "payment service overview",
         internal_to="Payment Service Overview")


# ---------------------------------------------------------------- Confluence

USERS = {"557058:aa11": "Jane Doe", "557058:bb22": "Sam Patel", "557058:cc33": "Alex Kim"}
SPACE = "NWP"


def page_link(title, text):
    return (f'<ac:link><ri:page ri:content-title="{title}" />'
            f"<ac:plain-text-link-body><![CDATA[{text}]]></ac:plain-text-link-body></ac:link>")


def mention(acc):
    return f'<ac:link><ri:user ri:account-id="{acc}" /></ac:link>'


def macro(name, body="", params=None, plain=False):
    ps = "".join(f'<ac:parameter ac:name="{k}">{v}</ac:parameter>' for k, v in (params or {}).items())
    inner = (f"<ac:plain-text-body><![CDATA[{body}]]></ac:plain-text-body>" if plain
             else (f"<ac:rich-text-body>{body}</ac:rich-text-body>" if body else ""))
    return f'<ac:structured-macro ac:name="{name}" ac:schema-version="1">{ps}{inner}</ac:structured-macro>'


PAGES = [
    {"id": "163841", "title": "Payments Platform", "parent": None, "labels": ["payments", "home"],
     "status": None, "restricted": None,
     "versions": [("557058:aa11", "2025-10-01T09:00:00Z", "Created space home"),
                  ("557058:bb22", "2026-06-12T10:00:00Z", "Added platform principles")],
     "body": [
         ("macro-toc", None, macro("toc")),
         ("heading", "About this space", "<h1>About this space</h1>"),
         ("macro-info", "This space is the source of truth for payments design.",
          macro("info", "<p>This space is the source of truth for payments design.</p>")),
         ("page-link", "Payment Service", "<p>Start with " + page_link("Payment Service", "the payment service") + ".</p>"),
         ("page-link", "Capture Runbook", "<p>On call? Read " + page_link("Capture Runbook", "the capture runbook") + ".</p>"),
         ("macro-details", "Payments Engineering",
          macro("details", "<table><tbody><tr><th>Owner</th><td>Payments Engineering</td></tr>"
                           "<tr><th>Review cycle</th><td>Quarterly</td></tr></tbody></table>")),
     ]},
    {"id": "163842", "title": "Payment Service", "parent": "163841", "labels": ["payments", "system"],
     "status": {"name": "verified", "by": "557058:cc33", "at": "2026-07-19T09:00:00Z"}, "restricted": None,
     "versions": [("557058:aa11", "2025-11-04T09:00:00Z", "First draft"),
                  ("557058:bb22", "2026-03-02T11:00:00Z", "Added capture flow"),
                  ("557058:aa11", "2026-07-18T09:12:00Z", "Event-driven v2 note")],
     "body": [
         ("heading", "Overview", "<h1>Overview</h1>"),
         ("text", "Authorises and captures customer payments.", "<p>Authorises and captures customer payments.</p>"),
         ("mention", "Jane Doe", "<p>Owner: " + mention("557058:aa11") + "</p>"),
         ("macro-warning", "Capture is synchronous until v2 ships.",
          macro("warning", "<p>Capture is synchronous until v2 ships.</p>")),
         ("macro-code", "SELECT status FROM captures WHERE created_at > now() - interval '1 hour';",
          macro("code", "SELECT status FROM captures WHERE created_at > now() - interval '1 hour';",
                {"language": "sql"}, plain=True)),
         ("macro-jira", "NWP-1234", macro("jira", params={"key": "NWP-1234", "server": "Example Jira"})),
         ("attachment", "payment-flow.png", '<ac:image><ri:attachment ri:filename="payment-flow.png" /></ac:image>'),
         ("page-link", "Capture Runbook", "<p>Operations: " + page_link("Capture Runbook", "capture runbook") + ".</p>"),
         ("table", "Europe", "<table><tbody><tr><th>Region</th><th>Peak TPS</th></tr>"
                             "<tr><td>Europe</td><td>1,200</td></tr></tbody></table>"),
     ]},
    {"id": "163843", "title": "Capture Runbook", "parent": "163841", "labels": ["runbook", "on-call"],
     "status": {"name": "verified", "by": "557058:aa11", "at": "2026-09-01T08:30:00Z"}, "restricted": None,
     "versions": [("557058:bb22", "2026-01-12T10:00:00Z", "Initial runbook"),
                  ("557058:bb22", "2026-08-30T14:20:00Z", "Added standby drain step")],
     "body": [
         ("heading", "When to use", "<h1>When to use</h1>"),
         ("text", "Use when capture latency rises above the alert threshold.",
          "<p>Use when capture latency rises above the alert threshold.</p>"),
         ("list-item", "Drain the capture queue to the standby worker",
          "<ol><li>Check the acquirer status page</li><li>Drain the capture queue to the standby worker</li></ol>"),
         ("macro-expand", "Restart the capture worker with the previous image tag.",
          macro("expand", "<p>Restart the capture worker with the previous image tag.</p>", {"title": "Rollback steps"})),
         ("macro-code", "kubectl rollout undo deployment/capture-worker",
          macro("code", "kubectl rollout undo deployment/capture-worker", {"language": "bash"}, plain=True)),
         ("mention", "Sam Patel", "<p>Escalate to " + mention("557058:bb22") + ".</p>"),
         ("page-link", "Payment Service", "<p>Background: " + page_link("Payment Service", "payment service") + ".</p>"),
         ("page-link", "Refund Policy", "<p>For refund disputes see " + page_link("Refund Policy", "refund policy") + ".</p>"),
     ]},
    {"id": "163844", "title": "Refund Policy", "parent": "163841", "labels": ["policy", "finance"],
     "status": None, "restricted": {"read": ["group:finance-leads"]},
     "versions": [("557058:cc33", "2026-05-02T13:00:00Z", "Policy draft")],
     "body": [
         ("heading", "Limits", "<h1>Limits</h1>"),
         ("macro-note", "Refunds over 500 GBP need a finance lead.",
          macro("note", "<p>Refunds over 500 GBP need a finance lead.</p>")),
     ]},
    {"id": "163845", "title": "Orders Table", "parent": "163842", "labels": ["data-asset", "orders"],
     "status": {"name": "rough draft"}, "restricted": None,
     "versions": [("557058:aa11", "2025-03-10T09:30:00Z", "Data dictionary"),
                  ("557058:aa11", "2025-08-14T10:00:00Z", "Added captured_amount")],
     "body": [
         ("heading", "Columns", "<h1>Columns</h1>"),
         ("table", "captured_amount", "<table><tbody><tr><th>Column</th><th>Type</th></tr>"
                                      "<tr><td>order_id</td><td>uuid</td></tr>"
                                      "<tr><td>captured_amount</td><td>numeric</td></tr></tbody></table>"),
         ("macro-details", "Data Platform", macro("details", "<table><tbody><tr><th>Owner</th><td>Data Platform</td></tr>"
                                                             "</tbody></table>")),
         ("macro-jira", "NWP-1290", macro("jira", params={"key": "NWP-1290", "server": "Example Jira"})),
         ("attachment", "orders-sample.csv",
          '<p>Sample: <ac:link><ri:attachment ri:filename="orders-sample.csv" /></ac:link></p>'),
         ("page-link", "Refund Policy", "<p>Refunds adjust captured_amount; see " + page_link("Refund Policy", "refunds") + ".</p>"),
     ]},
]


def confluence() -> None:
    out = OUT / "confluence"
    out.mkdir(parents=True, exist_ok=True)
    titles = {p["id"]: p["title"] for p in PAGES}
    for p in PAGES:
        xhtml = "".join(b[2] for b in p["body"])
        (out / f"{p['id']}.storage.xhtml").write_text(xhtml, encoding="utf-8")
        for kind, pr, _ in p["body"]:
            if kind == "page-link":
                probe("confluence", p["title"], "link-internal", pr, target=pr)
            elif kind != "macro-toc":
                probe("confluence", p["title"], kind, pr)
            else:
                probe("confluence", p["title"], kind, "", regenerable=True)
        last = p["versions"][-1]
        meta = {"id": p["id"], "type": "page", "title": p["title"], "space": {"key": SPACE, "name": "Northwind Pay"},
                "_links": {"webui": f"/spaces/{SPACE}/pages/{p['id']}/{p['title'].replace(' ', '+')}",
                           "base": WIKI},
                "version": {"number": len(p["versions"]), "by": {"accountId": last[0], "displayName": USERS[last[0]]},
                            "when": last[1], "message": last[2]},
                "history": {"createdBy": {"accountId": p["versions"][0][0], "displayName": USERS[p["versions"][0][0]]},
                            "createdDate": p["versions"][0][1],
                            "versions": [{"number": i + 1, "by": {"accountId": a, "displayName": USERS[a]},
                                          "when": w, "message": m} for i, (a, w, m) in enumerate(p["versions"])]},
                "ancestors": ([{"id": p["parent"], "title": titles[p["parent"]]}] if p["parent"] else []),
                "metadata": {"labels": {"results": [{"name": l} for l in p["labels"]]}},
                "contentState": p["status"], "restrictions": p["restricted"] or {}}
        (out / f"{p['id']}.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
        for field, val in (("title", p["title"]), ("page-id", p["id"]), ("space", SPACE),
                           ("created-by", USERS[p["versions"][0][0]]), ("last-modified-by", USERS[last[0]]),
                           ("last-modified", last[1]), ("version-number", str(len(p["versions"])))):
            probe("confluence", p["title"], "metadata", val, field=field)
        for l in p["labels"]:
            probe("confluence", p["title"], "metadata", l, field="label")
        if p["parent"]:
            probe("confluence", p["title"], "metadata", titles[p["parent"]], field="ancestor")
        if p["status"]:
            probe("confluence", p["title"], "metadata", p["status"]["name"], field="content-status")
        if p["restricted"]:
            probe("confluence", p["title"], "metadata", "restricted", field="restriction")
        for _, _, m in p["versions"]:
            probe("confluence", p["title"], "version", m)
    (out / "users.json").write_text(json.dumps(USERS, indent=2), encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for meta, build in zip(WORD, (build_overview, build_runbook, build_refund, build_orders, build_incident)):
        word_doc(meta, build)
    confluence()
    (OUT / "ground-truth.json").write_text(json.dumps(truth, indent=2), encoding="utf-8")
    print(f"corpus: {len(WORD)} Word documents, {len(PAGES)} Confluence pages, {len(truth)} probes -> {OUT}")


if __name__ == "__main__":
    main()
